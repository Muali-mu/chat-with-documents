"""
Each parser returns a list of (section_label, text) tuples so chunks can carry
a human-readable location (page number, slide number, heading) for citations.
"""
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader


def parse_pdf(path: Path) -> list[tuple[str, str]]:
    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            out.append((f"Page {i}", text))
    return out


def parse_docx(path: Path) -> list[tuple[str, str]]:
    doc = DocxDocument(str(path))
    out = []
    current_heading = "Document start"
    buffer = []

    def flush():
        if buffer:
            out.append((current_heading, "\n".join(buffer)))
            buffer.clear()

    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            flush()
            current_heading = para.text.strip() or current_heading
        elif para.text.strip():
            buffer.append(para.text)
    flush()
    return out


def parse_pptx(path: Path) -> list[tuple[str, str]]:
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            out.append((f"Slide {i}", "\n".join(texts)))
    return out


def parse_html(path: Path) -> list[tuple[str, str]]:
    soup = BeautifulSoup(path.read_text(errors="ignore"), "html.parser")
    text = soup.get_text(separator="\n")
    title = soup.title.string.strip() if soup.title and soup.title.string else path.name
    return [(title, text)]


def parse_txt(path: Path) -> list[tuple[str, str]]:
    return [(path.name, path.read_text(errors="ignore"))]


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".html": parse_html,
    ".htm": parse_html,
    ".txt": parse_txt,
    ".md": parse_txt,
}


def parse_file(path: Path) -> list[tuple[str, str]]:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return parser(path)
